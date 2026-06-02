from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(r"D:\QQ\命题二：外呼任务对话模型指令示例.xlsx")
OUTPUT_DIR = Path("examples/tasks")


def main() -> None:
    workbook = load_workbook(SOURCE, data_only=True)
    sheet = workbook.active
    rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[1]]
    task_specs = []
    for index, (_, content) in enumerate(rows[:2], start=1):
        text = str(content).replace("\r\n", "\n").replace("\r", "\n")
        sections = split_sections(text)
        if index == 1:
            spec = build_delivery_task(sections)
            filename = "fengmaotui_delivery_task.json"
        else:
            spec = build_course_task(sections)
            filename = "course_live_task.json"
        path = OUTPUT_DIR / filename
        path.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
        task_specs.append((filename, spec))

    # Keep the original demo path as a compatibility alias for tests and docs.
    (OUTPUT_DIR / "delivery_task.json").write_text(
        json.dumps(task_specs[0][1], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print("\n".join(f"{filename}: {spec['task_id']}" for filename, spec in task_specs))


def split_sections(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {}
    current = "body"
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        match = re.match(r"^#{1,3}\s*([^:#]+)(?::\s*(.*))?$", line)
        if match:
            current = normalize_heading(match.group(1))
            sections.setdefault(current, [])
            if match.group(2):
                sections[current].append(match.group(2).strip())
            continue
        sections.setdefault(current, []).append(line)
    return {key: "\n".join(value).strip() for key, value in sections.items()}


def normalize_heading(value: str) -> str:
    heading = value.strip().lower()
    if "role" in heading:
        return "role"
    if "task" in heading:
        return "task"
    if "opening" in heading:
        return "opening_line"
    if "call flow" in heading or "conversation flow" in heading:
        return "flow"
    if "knowledge" in heading or "faq" in heading:
        return "faq"
    if "constraint" in heading:
        return "constraints"
    return heading


def clean_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = text.replace("`", "")
    return text.strip(" -\t")


def bullet_lines(text: str) -> list[str]:
    lines = []
    for line in text.splitlines():
        cleaned = clean_markdown(re.sub(r"^\d+[\.\)]\s*", "", line))
        if cleaned:
            lines.append(cleaned)
    return lines


def flow_steps(text: str, fallback: list[str]) -> list[dict]:
    lines = bullet_lines(text)
    candidates = [
        line
        for line in lines
        if not line.startswith("#") and len(line) > 4 and not line.startswith("参考")
    ]
    if not candidates:
        candidates = fallback
    return [
        {"step_id": f"step_{index:02d}", "description": line, "required": True}
        for index, line in enumerate(candidates[:12], start=1)
    ]


def faq_items(text: str, fallback: list[dict]) -> list[dict]:
    lines = bullet_lines(text)
    items = []
    for line in lines:
        if len(line) < 6:
            continue
        question = line
        answer = line
        if "：" in line:
            question, answer = line.split("：", 1)
        elif ":" in line:
            question, answer = line.split(":", 1)
        items.append({"question": question.strip(), "answer": answer.strip()})
    return items or fallback


def constraints(text: str, max_reply_chars: int) -> dict:
    terms = ["保证收益", "一定返钱", "优惠券", "折扣券", "内部政策"]
    return {
        "max_reply_chars": max_reply_chars,
        "tone": "电话口语、礼貌、自然、简短",
        "forbidden_terms": terms,
        "privacy_fields": ["身份证号", "手机号", "地址"],
    }


def build_delivery_task(sections: dict[str, str]) -> dict:
    return {
        "task_id": "fengmaotui_delivery_task",
        "role": clean_markdown(sections.get("role", "履约数字人")),
        "task": clean_markdown(
            sections.get("task", "通知骑手飞毛腿合同已成功签约，并提醒开始配送。")
        ),
        "opening_line": clean_markdown(
            sections.get(
                "opening_line",
                "您好，请问是${rider_name}吗？我是站长，来电通知您的飞毛腿合同已生效。",
            )
        ),
        "flow_steps": flow_steps(
            sections.get("flow", ""),
            ["通知合同生效", "询问是否开始配送", "说明配送要求", "处理异议并礼貌结束"],
        ),
        "faq": faq_items(
            sections.get("faq", ""),
            [
                {
                    "question": "如何退出飞毛腿",
                    "answer": "需要前一天指定时间前在 App 的飞毛腿报名入口取消。",
                }
            ],
        ),
        "constraints": constraints(sections.get("constraints", ""), 30),
        "tools": default_tools(),
    }


def build_course_task(sections: dict[str, str]) -> dict:
    return {
        "task_id": "course_live_task",
        "role": clean_markdown(
            sections.get("role", "Customer Support Specialist for Course Publishing Platform")
        ),
        "task": clean_markdown(
            sections.get(
                "task",
                "通知课程发布客户页面新增标准直播和低延迟直播选项，并确认客户是否知晓和会使用。",
            )
        ),
        "opening_line": clean_markdown(
            sections.get("opening_line", "您好，请问您是课程发布或校验系统的负责人吗？")
        ),
        "flow_steps": flow_steps(
            sections.get("flow", ""),
            [
                "确认是否为负责人",
                "说明新增标准直播和低延迟直播选项",
                "确认客户是否知晓",
                "说明两种直播模式差异",
                "确认前端是否可见并指导配置",
                "提醒企业微信通知",
                "礼貌结束",
            ],
        ),
        "faq": faq_items(
            sections.get("faq", ""),
            [
                {
                    "question": "标准直播和低延迟直播有什么区别",
                    "answer": "标准直播成本较低、延迟约 5-10 秒；低延迟直播约 1-2 秒，适合互动课程。",
                },
                {
                    "question": "低延迟直播是否更贵",
                    "answer": "低延迟直播链路保障更强，成本通常更高。",
                },
            ],
        ),
        "constraints": constraints(sections.get("constraints", ""), 20),
        "tools": default_tools(),
    }


def default_tools() -> list[str]:
    return [
        "transfer_to_human",
        "query_faq",
        "record_rejection",
        "schedule_callback",
        "create_ticket",
        "update_task_status",
    ]


if __name__ == "__main__":
    main()
